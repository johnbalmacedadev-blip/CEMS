#!/usr/bin/env python3
"""Export PURCHASE PRICES + financial totals from PRIVATE PURCHASE PRICES workbook."""

from __future__ import annotations

import json
import os
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("openpyxl is required", file=sys.stderr)
    sys.exit(1)


def norm_plate(value) -> str:
    if value is None:
        return ""
    return re.sub(r"[^A-Z0-9]", "", str(value).upper())


def to_float(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).replace(",", "").replace("₱", "").strip()
    try:
        return float(s)
    except ValueError:
        return None


def main() -> int:
    src = os.environ.get("PURCHASE_PRICES_SRC") or str(
        Path(__file__).resolve().parents[1] / "storage" / "app" / "PRIVATE_PURCHASE_PRICES.xlsx"
    )
    out = os.environ.get("PURCHASE_PRICES_OUT") or str(
        Path(__file__).resolve().parents[1] / "storage" / "app" / "purchase_prices_import.json"
    )

    if not Path(src).is_file():
        print(f"File not found: {src}", file=sys.stderr)
        return 1

    wb = openpyxl.load_workbook(src, data_only=True, read_only=True)

    prices: dict[str, float] = {}
    if "PURCHASE PRICES" in wb.sheetnames:
        ws = wb["PURCHASE PRICES"]
        for row in ws.iter_rows(min_row=2, max_col=2, values_only=True):
            plate = norm_plate(row[0] if row else None)
            amount = to_float(row[1] if row and len(row) > 1 else None)
            if plate and amount is not None and amount >= 0:
                prices[plate] = round(amount, 2)

    # Later sheets overwrite earlier ones so Released wins over Reserved/Available.
    tabs = [
        ("FLAGSHIP - AVAILABLE", "Available", "Flagship"),
        ("ANNEX - AVAILABLE", "Available", "Annex"),
        ("FLAGSHIP - RESERVED", "Reserved", "Flagship"),
        ("ANNEX - RESERVED", "Reserved", "Annex"),
        ("FLAGSHIP RELEASED", "Released", "Flagship"),
        ("FLAGSHIP - RELEASED", "Released", "Flagship"),
        ("ANNEX - RELEASED", "Released", "Annex"),
    ]

    financials: dict[str, dict] = {}
    for name, status, branch in tabs:
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        for row in ws.iter_rows(min_row=2, max_col=21, values_only=True):
            plate = norm_plate(row[15] if row and len(row) > 15 else None)
            if not plate or len(plate) < 3:
                continue
            purchase = to_float(row[0] if row else None)
            costs = to_float(row[1] if row and len(row) > 1 else None)
            revenue = to_float(row[2] if row and len(row) > 2 else None)
            profit = to_float(row[3] if row and len(row) > 3 else None)
            if purchase is None and costs is None and revenue is None and profit is None:
                continue
            if (purchase or 0) == 0 and (costs or 0) == 0 and (revenue or 0) == 0 and (profit or 0) == 0:
                continue
            financials[plate] = {
                "plate": plate,
                "purchase_price": purchase,
                "total_costs": costs,
                "total_revenue": revenue,
                "total_profit": profit,
                "status": status,
                "branch": branch,
                "sheet": name,
            }

    wb.close()

    payload = {
        "source": Path(src).name,
        "purchase_prices": [{"plate": p, "purchase_price": v} for p, v in sorted(prices.items())],
        "financials": list(financials.values()),
    }
    Path(out).parent.mkdir(parents=True, exist_ok=True)
    Path(out).write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    print(
        f"Wrote {out} | prices={len(prices)} financials={len(financials)}",
        flush=True,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
